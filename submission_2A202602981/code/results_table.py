"""results_table.py — Quản lý kết quả và xuất ra file Excel.

Lưu kết quả từng lần chạy ra JSON, sau đó xuất ra experiments.xlsx theo mẫu.
"""
from __future__ import annotations

import json
import os
from pathlib import Path


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] ra JSON."""
    os.makedirs(results_dir, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    file_path = os.path.join(results_dir, f"{exp_id}.json")

    data_to_save = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"]
    }

    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)

    return file_path


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict."""
    if not os.path.exists(results_dir):
        return []

    results = []
    p = Path(results_dir)
    for json_file in sorted(p.glob("*.json")):
        with open(json_file, "r", encoding="utf-8") as f:
            results.append(json.load(f))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng Experiments."""
    cfg = result["cfg"]
    summary = result["summary"]
    exp_id = cfg["exp_id"]

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce"),
        "optimizer": cfg.get("optimizer", "sgd_momentum"),
        "lr": cfg.get("lr", 0.0),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": str(cfg.get("hidden", (256, 128))),
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": cfg.get("clip_norm", ""),
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss", 0.0),
        "best_val_loss": summary.get("best_val_loss", 0.0),
        "best_epoch": summary.get("best_epoch", 1),
        "final_train_loss": summary.get("final_train_loss", 0.0),
        "final_val_loss": summary.get("final_val_loss", 0.0),
        "val_acc": summary.get("val_acc", 0.0),
        "val_macro_f1": summary.get("val_macro_f1", 0.0),
        "time_per_epoch_s": summary.get("time_per_epoch_s", 0.0),
        "peak_mem_MB": summary.get("peak_mem_MB", 0.0),
        "diverged": summary.get("diverged", False),
        "eval_acc": eval_scores.get("acc", "") if eval_scores else "",
        "eval_macro_f1": eval_scores.get("macro_f1", "") if eval_scores else "",
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet 'Experiments' của mẫu Excel."""
    try:
        import openpyxl
    except ImportError:
        print("CẢNH BÁO: Chưa cài đặt thư viện 'openpyxl'. Vui lòng chạy `pip install openpyxl` để xuất file Excel.")
        return

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    # Đọc tiêu đề dòng 1
    col_mapping = {}
    for col_idx in range(1, ws.max_column + 1):
        header_val = ws.cell(row=1, column=col_idx).value
        if header_val:
            col_mapping[str(header_val).strip()] = col_idx

    formula_cols = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}

    current_row = 2
    for r in rows:
        for key, val in r.items():
            if key in col_mapping and key not in formula_cols:
                col_num = col_mapping[key]
                ws.cell(row=current_row, column=col_num, value=val)
        current_row += 1

    wb.save(out_path)
    print(f"Đã lưu thành công bảng kết quả tại: {out_path}")
